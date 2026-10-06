"""Isolamento de escopo — o dado de uma criança só sai para quem pode vê-la.

Fase 3A (segurança). Mundo SINTÉTICO, com nomes impossíveis de confundir com
dado real, para que a simples presença de um nome numa resposta seja a prova
do vazamento — sem depender do formato do JSON, do CSV, do XLSX ou do PDF.

    ESCOLA_X (rede R1, ano ativo 2026)
      "5º Ano B" 2025 — Professora A   → ALUNO_PROFESSOR_A_2025
      "4º Ano A" 2026 — Professora A   → ALUNO_PROFESSOR_A_2026
      "5º Ano B" 2026 — Professora B   → ALUNO_PROFESSOR_B_2026
    ESCOLA_Y (sem rede, ano ativo 2026)
      "4º Ano A" 2026 — Professor C    → ALUNO_OUTRA_ESCOLA

O caso que motivou a suíte: a turma "5º Ano B" existe em 2025 (da A) e em 2026
(da B). A exportação recortava o professor comparando o NOME da turma, e a
lista de turmas permitidas trazia turmas de qualquer ano — a A exportava as
crianças da B. Ver `permissoes.turmas_permitidas`.
"""
from io import BytesIO

import pytest
from fastapi.testclient import TestClient

from app.core.config import settings
from app.core.security import hash_senha
from app.main import app
from app.models import (Aluno, Escola, Livro, Matricula, Nota, Professor, Rede,
                        Turma, Usuario)

A25, A26, B26, OUTRA = ("ALUNO_PROFESSOR_A_2025", "ALUNO_PROFESSOR_A_2026",
                        "ALUNO_PROFESSOR_B_2026", "ALUNO_OUTRA_ESCOLA")
SENHA = "s3nh4-sintetica"


class Mundo:
    def __init__(self, db):
        self.db = db
        self.ids: dict[str, int] = {}
        self._clientes: dict[str, TestClient] = {}

    def cli(self, quem: str) -> TestClient:
        if quem not in self._clientes:
            c = TestClient(app)
            r = c.post("/api/v1/auth/login",
                       data={"username": f"{quem}@sint.local", "password": SENHA})
            assert r.status_code == 200, (quem, r.text)
            c.headers["Authorization"] = f"Bearer {r.json()['access_token']}"
            self._clientes[quem] = c
        return self._clientes[quem]


@pytest.fixture()
def mundo(db, monkeypatch, tmp_path):
    # As exportações gravam uma cópia em EXPORTS_DIR; o teste não suja o repo.
    monkeypatch.setattr(settings, "EXPORTS_DIR", tmp_path)
    m = Mundo(db)
    r1, r2 = Rede(nome="REDE_SINT_1"), Rede(nome="REDE_SINT_2")
    db.add_all([r1, r2]); db.flush()
    x = Escola(nome="ESCOLA_X_SINTETICA", ano_letivo_ativo=2026, rede_id=r1.id)
    y = Escola(nome="ESCOLA_Y_SINTETICA", ano_letivo_ativo=2026)
    db.add_all([x, y]); db.flush()
    m.ids.update(x=x.id, y=y.id, r1=r1.id, r2=r2.id)

    def usuario(apelido, escola, cargo, **kw):
        u = Usuario(escola_id=escola, nome=apelido.upper(), email=f"{apelido}@sint.local",
                    senha_hash=hash_senha(SENHA), cargo=cargo, **kw)
        db.add(u); db.flush()
        m.ids["u_" + apelido] = u.id
        return u

    usuario("coord_x", x.id, "coordenador")
    usuario("admin_x", x.id, "admin")
    usuario("prof_a", x.id, "professor")
    usuario("prof_b", x.id, "professor")
    usuario("coord_y", y.id, "coordenador")
    usuario("sec_r1", None, "coordenador", rede_id=r1.id)
    usuario("sec_r2", None, "coordenador", rede_id=r2.id)
    usuario("global", x.id, "admin", is_global=True)

    pa = Professor(escola_id=x.id, nome="PROF_A_SINT", email="prof_a@sint.local")
    pb = Professor(escola_id=x.id, nome="PROF_B_SINT", email="prof_b@sint.local")
    pc = Professor(escola_id=y.id, nome="PROF_C_Y_SINT", email="prof_c@sint.local")
    db.add_all([pa, pb, pc]); db.flush()

    def turma(chave, escola, nome, ano, prof):
        # "5º Ano B" -> série "5º Ano"
        t = Turma(escola_id=escola, nome=nome, ano_escolar=nome.rsplit(" ", 1)[0],
                  ano_letivo=ano, professor_id=prof.id)
        db.add(t); db.flush()
        m.ids[chave] = t.id
        return t

    t_a25 = turma("t_a25", x.id, "5º Ano B", 2025, pa)
    t_a26 = turma("t_a26", x.id, "4º Ano A", 2026, pa)
    t_b26 = turma("t_b26", x.id, "5º Ano B", 2026, pb)
    t_y = turma("t_y", y.id, "4º Ano A", 2026, pc)

    def aluno(nome, escola, t, ano, posicao):
        a = Aluno(escola_id=escola, nome=nome)
        db.add(a); db.flush()
        db.add(Matricula(escola_id=escola, aluno_id=a.id, turma_id=t.id, ano_letivo=ano))
        db.add(Nota(escola_id=escola, aluno_id=a.id, ano_letivo=ano,
                    nota_elefante=70.0 + posicao, nota_matific=60.0 + posicao,
                    nota_geral=65.0 + posicao, posicao=posicao,
                    aferido_leitura=True, aferido_matematica=True,
                    posicao_leitura=posicao, posicao_matematica=posicao))
        m.ids[nome] = a.id
        return a

    aluno(A25, x.id, t_a25, 2025, 1)
    aluno(A26, x.id, t_a26, 2026, 2)
    aluno(B26, x.id, t_b26, 2026, 1)
    aluno(OUTRA, y.id, t_y, 2026, 1)
    livro_y = Livro(escola_id=y.id, titulo="LIVRO_SINT_ESCOLA_Y", nivel_codigo="A")
    db.add(livro_y); db.flush()
    m.ids["livro_y"] = livro_y.id
    db.commit()
    return m


def conteudo(resposta, formato: str) -> str:
    """O texto que a pessoa leria no arquivo — para procurar nomes nele."""
    if formato == "csv":
        return resposta.content.decode("utf-8-sig")
    if formato == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(resposta.content))
        return " ".join(str(c.value) for ws in wb.worksheets
                        for linha in ws.iter_rows() for c in linha if c.value is not None)
    if formato == "pdf":
        from pypdf import PdfReader
        return " ".join((p.extract_text() or "") for p in PdfReader(BytesIO(resposta.content)).pages)
    return resposta.text


# ===========================================================================
# 1) EXPORTAÇÕES — o vazamento reproduzido na Fase 2.1
# ===========================================================================
TIPOS = ["alunos", "ranking", "ranking_leitura", "ranking_matematica"]
FORMATOS = ["csv", "xlsx", "pdf"]


@pytest.mark.parametrize("formato", FORMATOS)
@pytest.mark.parametrize("tipo", TIPOS)
def test_professora_exporta_so_as_criancas_dela(mundo, tipo, formato):
    r = mundo.cli("prof_a").get(
        f"/api/v1/escolas/{mundo.ids['x']}/relatorios/{tipo}?formato={formato}")
    assert r.status_code == 200, r.text
    texto = conteudo(r, formato)
    # Controle positivo: sem ele, "não achei o nome alheio" provaria nada
    # (um arquivo vazio, ou um PDF sem texto extraível, também não acha).
    assert A26 in texto, f"{tipo}/{formato}: a criança DELA não saiu — o teste não prova nada"
    vazados = [n for n in (B26, OUTRA, A25) if n in texto]
    assert vazados == [], f"{tipo}/{formato} vazou: {vazados}"


@pytest.mark.parametrize("tipo", TIPOS)
def test_a_outra_professora_tambem_so_ve_as_dela(mundo, tipo):
    r = mundo.cli("prof_b").get(
        f"/api/v1/escolas/{mundo.ids['x']}/relatorios/{tipo}?formato=csv")
    assert r.status_code == 200
    texto = conteudo(r, "csv")
    assert B26 in texto
    assert A26 not in texto and OUTRA not in texto


def test_coordenadora_exporta_a_escola_inteira_e_so_ela(mundo):
    texto = conteudo(mundo.cli("coord_x").get(
        f"/api/v1/escolas/{mundo.ids['x']}/relatorios/alunos?formato=csv"), "csv")
    assert A26 in texto and B26 in texto      # a escola dela, inteira
    assert OUTRA not in texto                 # e nada de outra escola


@pytest.mark.parametrize("quem", ["sec_r1", "coord_y", "sec_r2"])
def test_quem_nao_e_da_escola_nao_exporta(mundo, quem):
    r = mundo.cli(quem).get(
        f"/api/v1/escolas/{mundo.ids['x']}/relatorios/alunos?formato=csv")
    assert r.status_code in (403, 404)
    assert A26 not in r.text and B26 not in r.text


# ===========================================================================
# 2) RELATÓRIO POR PERÍODO (rotas próprias, id de turma/aluno na query)
# ===========================================================================
def test_relatorio_por_periodo_da_professora(mundo):
    base = f"/api/v1/escolas/{mundo.ids['x']}/relatorios/periodo?periodo=ano_letivo"
    cli = mundo.cli("prof_a")
    r = cli.get(base)
    assert r.status_code == 200
    assert B26 not in r.text and OUTRA not in r.text

    # turma da outra professora, mesmo nome da antiga dela
    assert cli.get(base + f"&escopo=turma&turma_id={mundo.ids['t_b26']}").status_code == 404
    # aluno da outra professora
    assert cli.get(base + f"&escopo=aluno&aluno_id={mundo.ids[B26]}").status_code == 404
    # turma de outra escola
    assert cli.get(base + f"&escopo=turma&turma_id={mundo.ids['t_y']}").status_code == 404


def test_turma_de_ano_anterior_nao_e_mais_escopo_do_professor(mundo):
    """Mudança deliberada da 3A: a designação que vale é a do ano ATIVO.
    A turma de 2025 da professora não abre mais por id — e não havia tela
    dela que usasse isso (as telas de turma são de gestão)."""
    base = f"/api/v1/escolas/{mundo.ids['x']}"
    cli = mundo.cli("prof_a")
    assert cli.get(f"{base}/turmas/{mundo.ids['t_a25']}/resumo").status_code == 404
    assert cli.get(f"{base}/relatorios/periodo?periodo=ano_letivo&escopo=turma"
                   f"&turma_id={mundo.ids['t_a25']}").status_code == 404
    # a turma ATUAL dela continua abrindo
    assert cli.get(f"{base}/turmas/{mundo.ids['t_a26']}/resumo").status_code == 200


def test_relatorio_por_periodo_em_pdf_respeita_o_mesmo_recorte(mundo):
    r = mundo.cli("prof_a").get(
        f"/api/v1/escolas/{mundo.ids['x']}/relatorios/periodo.pdf?periodo=ano_letivo")
    assert r.status_code == 200
    texto = conteudo(r, "pdf")
    assert B26 not in texto and OUTRA not in texto


# ===========================================================================
# 3) TELAS E LISTAS DO PROFESSOR
# ===========================================================================
@pytest.mark.parametrize("caminho", [
    "/alunos?por_pagina=100", "/ranking", "/ranking?dimensao=leitura",
    "/ranking?dimensao=matematica", "/ranking/leitura?periodo=ano_letivo",
    "/ranking/matematica?periodo=ano_letivo", "/ranking-evolucao?periodo=ano_letivo",
    "/premiacoes?periodo=tudo", "/insights", "/gamificacao/ranking-xp",
    "/pesquisa?q=ALUNO", "/nao-aferidos", "/elefante", "/matific", "/dashboard",
])
def test_professora_nao_recebe_crianca_de_fora(mundo, caminho):
    r = mundo.cli("prof_a").get(f"/api/v1/escolas/{mundo.ids['x']}{caminho}")
    assert r.status_code == 200, (caminho, r.status_code, r.text[:200])
    vazados = [n for n in (B26, OUTRA, A25) if n in r.text]
    assert vazados == [], f"{caminho} vazou: {vazados}"


@pytest.mark.parametrize("sufixo", ["perfil", "evolucao?dias=30", "evolucao-leitura",
                                    "leituras", "linha-do-tempo", "espelho"])
def test_ficha_de_crianca_alheia_nao_abre(mundo, sufixo):
    r = mundo.cli("prof_a").get(
        f"/api/v1/escolas/{mundo.ids['x']}/alunos/{mundo.ids[B26]}/{sufixo}")
    assert r.status_code in (403, 404)
    assert B26 not in r.text


def test_certificado_de_crianca_alheia_nao_sai(mundo):
    r = mundo.cli("prof_a").get(
        f"/api/v1/escolas/{mundo.ids['x']}/certificados/{mundo.ids[B26]}")
    assert r.status_code in (403, 404)


# ===========================================================================
# 4) SECRETARIA — rede, nunca criança por criança
# ===========================================================================
@pytest.mark.parametrize("caminho", [
    "/alunos?por_pagina=100", "/ranking", "/insights", "/premiacoes?periodo=tudo",
    "/gamificacao/ranking-xp", "/pesquisa?q=ALUNO", "/elefante", "/matific",
])
def test_secretaria_da_rede_nao_recebe_nome_de_crianca(mundo, caminho):
    r = mundo.cli("sec_r1").get(f"/api/v1/escolas/{mundo.ids['x']}{caminho}")
    assert r.status_code in (200, 403, 404, 422)
    vazados = [n for n in (A26, B26, A25, OUTRA) if n in r.text]
    assert vazados == [], f"{caminho} entregou nome de criança à Secretaria: {vazados}"


def test_secretaria_de_outra_rede_nao_entra_na_escola(mundo):
    r = mundo.cli("sec_r2").get(f"/api/v1/escolas/{mundo.ids['x']}/resumo-escola")
    assert r.status_code in (403, 404)


def test_secretaria_ve_o_agregado_da_propria_rede(mundo):
    r = mundo.cli("sec_r1").get(f"/api/v1/redes/{mundo.ids['r1']}/dashboard")
    assert r.status_code == 200
    assert "ESCOLA_X_SINTETICA" in r.text
    assert all(n not in r.text for n in (A26, B26, OUTRA))


# ===========================================================================
# 5) IDOR — a escola certa na URL e o id de OUTRA escola no caminho
# ===========================================================================
# A varredura da Fase 2.1 trocou o id da ESCOLA na URL; esta troca o id da
# ENTIDADE e mantém a escola do usuário. `exigir_turma_permitida` e
# `exigir_aluno_permitido` retornam cedo para a gestão — então quem garante que
# a turma/aluno é DA escola é cada rota. Aqui isso é testado em todas as rotas
# GET que levam id de entidade no caminho, geradas a partir do próprio app:
# rota nova entra no teste sozinha.
def _rotas_com_entidade():
    import re
    saida = []
    for r in app.routes:
        if not hasattr(r, "methods") or "GET" not in r.methods:
            continue
        if "/escolas/{escola_id}/" not in r.path:
            continue
        params = [p for p in re.findall(r"\{(\w+)\}", r.path) if p != "escola_id"]
        if any(p in ("aluno_id", "turma_id", "usuario_id") for p in params):
            saida.append(r.path)
    return sorted(set(saida))


ALHEIOS_Y = ("ALUNO_OUTRA_ESCOLA", "ESCOLA_Y_SINTETICA", "PROF_C_Y_SINT",
             "coord_y@sint.local", "LIVRO_SINT_ESCOLA_Y")


@pytest.mark.parametrize("modelo", _rotas_com_entidade())
def test_coordenadora_nao_abre_entidade_de_outra_escola(mundo, modelo):
    caminho = (modelo
               .replace("{escola_id}", str(mundo.ids["x"]))
               .replace("{aluno_id}", str(mundo.ids[OUTRA]))
               .replace("{turma_id}", str(mundo.ids["t_y"]))
               .replace("{usuario_id}", str(mundo.ids["u_coord_y"]))
               .replace("{plataforma}", "elefante"))
    r = mundo.cli("coord_x").get(caminho)
    assert r.status_code < 500, f"{modelo}: {r.status_code} — rota quebrou com id alheio"
    vazados = [n for n in ALHEIOS_Y if n in r.text]
    assert vazados == [], f"{modelo} ({r.status_code}) entregou dado de outra escola: {vazados}"


def test_a_varredura_de_idor_cobre_as_rotas_que_importam():
    """Se a introspecção quebrar e devolver pouca coisa, a varredura viraria um
    teste verde vazio. Trava um piso e as rotas-chave."""
    rotas = _rotas_com_entidade()
    assert len(rotas) >= 12
    for chave in ("/alunos/{aluno_id}/perfil", "/turmas/{turma_id}/resumo",
                  "/certificados/{aluno_id}", "/usuarios/{usuario_id}/turmas"):
        assert any(chave in r for r in rotas), chave


@pytest.mark.parametrize("tipo,chave", [("aluno", OUTRA), ("turma", "t_y")])
def test_comparador_nao_compara_com_entidade_de_outra_escola(mundo, tipo, chave):
    x = mundo.ids["x"]
    proprio = mundo.ids[A26] if tipo == "aluno" else mundo.ids["t_a26"]
    r = mundo.cli("coord_x").get(
        f"/api/v1/escolas/{x}/comparar?tipo_a={tipo}&id_a={proprio}"
        f"&tipo_b={tipo}&id_b={mundo.ids[chave]}")
    assert r.status_code in (403, 404), r.text[:200]
    assert all(n not in r.text for n in ALHEIOS_Y)
